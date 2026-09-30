"""Experimento C (ADR-0090): acuerdo entre la calificación humana y la segunda opinión automática.

    python comparar_opiniones.py c1

Solo se corre cuando la hoja de evaluación está completa. Genera datos/acuerdo-<corrida>.json con:
- acuerdo y kappa de Cohen con las cuatro categorías de la rúbrica y con dos (relevante / no relevante, en sentido amplio);
- la matriz de confusión (persona en filas, modelo en columnas);
- como análisis secundario, el resultado de H1 con la calificación automática, con las mismas pruebas que analizar.py."""

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from analizar import MAL_SUSTENTADA, RELEVANTE, signos, wilcoxon

BASE = Path(__file__).parent
RUBRICA = ["relevante y correcta", "relevante pero mal sustentada", "ruido", "falsa"]


def kappa(pares: list[tuple[str, str]]) -> float:
    n = len(pares)
    observado = sum(a == b for a, b in pares) / n
    fa, fb = Counter(a for a, _ in pares), Counter(b for _, b in pares)
    esperado = sum(fa[c] * fb[c] for c in set(fa) | set(fb)) / (n * n)
    return round((observado - esperado) / (1 - esperado), 3) if esperado < 1 else 1.0


def main(corrida: str) -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    clave = json.loads((BASE / "clave" / f"clave-{corrida}.json").read_text(encoding="utf-8"))["observaciones"]
    opinion = json.loads((BASE / "segunda-opinion" / f"opinion-{corrida}.json").read_text(encoding="utf-8"))
    modelo = {c["id"]: c["calificacion"] for pr in opinion["prs"] for c in pr["calificaciones"]}
    ws = load_workbook(BASE / "evaluacion-ciega" / "hoja-de-evaluacion.xlsx")["Calificación"]
    persona = {f[0].value: f[6].value for f in ws.iter_rows(min_row=2)}
    faltan = [i for i in clave if persona.get(i) not in RUBRICA or modelo.get(i) not in RUBRICA]
    if faltan:
        sys.exit(f"Faltan {len(faltan)} calificaciones (p. ej. {faltan[:5]}).")

    pares = [(persona[i], modelo[i]) for i in clave]

    def amplia(c: str) -> str:
        return "relevante" if c in (RELEVANTE, MAL_SUSTENTADA) else "no relevante"

    pares2 = [(amplia(a), amplia(b)) for a, b in pares]
    confusion = {a: {b: sum(1 for x, y in pares if x == a and y == b) for b in RUBRICA} for a in RUBRICA}
    por_pr: dict[str, dict[str, int]] = defaultdict(lambda: {"contexto": 0, "solo_diff": 0})
    por_cond: dict[str, Counter] = {"contexto": Counter(), "solo_diff": Counter()}
    for i, k in clave.items():
        por_cond[k["condicion"]][modelo[i]] += 1
        por_pr[i.split("-")[0]][k["condicion"]] += amplia(modelo[i]) == "relevante"
    dif = [v["contexto"] - v["solo_diff"] for v in por_pr.values()]
    resultado = {
        "corrida": corrida, "observaciones": len(pares), "modelo_calificador": opinion["modelos"],
        "cuatro_categorias": {"acuerdo": round(sum(a == b for a, b in pares) / len(pares), 3), "kappa": kappa(pares)},
        "dos_categorias": {"acuerdo": round(sum(a == b for a, b in pares2) / len(pares2), 3), "kappa": kappa(pares2)},
        "confusion_persona_filas_modelo_columnas": confusion,
        "h1_con_la_calificacion_automatica": {
            c: {"observaciones": sum(v.values()), "calificaciones": dict(v),
                "proporcion_relevantes_amplia": round((v[RELEVANTE] + v[MAL_SUSTENTADA]) / sum(v.values()), 3)}
            for c, v in por_cond.items()
        } | {"prueba_pareada_amplia": {"wilcoxon": wilcoxon(dif), "signos": signos(dif)}},
    }
    (BASE / "datos" / f"acuerdo-{corrida}.json").write_text(json.dumps(resultado, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps(resultado, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "c1")
