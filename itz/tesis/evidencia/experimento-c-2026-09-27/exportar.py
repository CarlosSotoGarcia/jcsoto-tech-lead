"""Experimento C (ADR-0089): exporta de Mongo los resultados de una corrida y arma el paquete de evaluación ciega.

    python exportar.py c1

Genera:
- datos/experimento_c-<corrida>.json   resultados crudos (las dos condiciones, con fuente y severidad)
- datos/metricas-<corrida>.json        llamadas al modelo (duración, costo nocional, modelo informado por el CLI)
- datos/resumen-<corrida>.json         conteos por condición antes de calificar
- evaluacion-ciega/hoja-de-evaluacion.xlsx   observaciones mezcladas y anónimas, con la rúbrica como lista
- evaluacion-ciega/contexto/Pnn.md           lo que quien califica necesita para juzgar (paquete, HU, casos, enlace al cambio)
- clave/clave-<corrida>.json            une cada identificador anónimo con su condición (fuera de git hasta calificar)

La clave no se versiona hasta que termine la calificación; su SHA-256 se anota en el README para demostrar después que no cambió."""

import hashlib
import json
import random
import subprocess
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pymongo import MongoClient

BASE = Path(__file__).parent
TARGET = Path(__file__).resolve().parents[4] / "loom" / "loom_target"  # repositorios de control de Loom
SEMILLA_CIEGA = 7319  # distinta de la del orden de las condiciones
RUBRICA = ["relevante y correcta", "relevante pero mal sustentada", "ruido", "falsa"]
PROYECTOS = {"9c764535fb44470a8dc92fc74c8b835e": "E1c", "b17137fd304348d7a48532ba5db89270": "E3c"}


def gh(*args: str) -> str:
    r = subprocess.run(["gh", "api", *args], capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        raise RuntimeError(r.stderr[:300])
    return r.stdout


def enlace_del_cambio(url_pr: str) -> str:
    """URL del cambio que vio la primera revisión (primer commit del PR), el mismo diff de las dos condiciones."""
    dueno_repo, numero = url_pr.split("github.com/")[1].split("/pull/")
    primero = json.loads(gh(f"repos/{dueno_repo}/pulls/{numero}/commits"))[0]
    return f"https://github.com/{dueno_repo}/commit/{primero['sha']}"


def md_de_la_primera_revision(proyecto_id: str, referencia: str, grupo: str, pt: str) -> str:
    """Mismo criterio que el runner (loom_backend.experimentos.revision_contexto): el archivo del repositorio de control en
    el commit anterior a «<referencia>: revisión 1»."""
    repo = TARGET / proyecto_id
    log = subprocess.run(["git", "log", "--format=%H%x09%s"], cwd=repo, capture_output=True, text=True, encoding="utf-8", check=True)
    for linea in log.stdout.splitlines():
        sha, asunto = linea.split("\t", 1)
        if asunto.startswith(f"{referencia}: revisión 1"):
            return subprocess.run(["git", "show", f"{sha}^:{grupo}/paquetes/{pt}.md"], cwd=repo, capture_output=True, text=True,
                                  encoding="utf-8", check=True).stdout
    raise LookupError(f"{referencia}: sin commit de su primera revisión")


def main(corrida: str) -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    db = MongoClient("mongodb://localhost:27017")["loom"]
    docs = list(db.experimento_c.find({"corrida": corrida}, {"_id": 0}).sort([("proyecto_id", 1), ("referencia", 1), ("condicion", 1)]))
    if not docs:
        sys.exit(f"No hay resultados de la corrida {corrida}.")
    pares: dict[tuple[str, str], dict[str, dict]] = {}
    for d in docs:
        pares.setdefault((d["proyecto_id"], d["referencia"]), {})[d["condicion"]] = d
    incompletos = [k for k, v in pares.items() if len(v) != 2]
    if incompletos:
        print("AVISO: pares incompletos (se excluyen):", incompletos)
        pares = {k: v for k, v in pares.items() if len(v) == 2}

    (BASE / "datos").mkdir(exist_ok=True)
    (BASE / "datos" / f"experimento_c-{corrida}.json").write_text(json.dumps(docs, ensure_ascii=False, indent=1), encoding="utf-8")
    metricas = list(db.metricas.find(
        {"skill": {"$regex": "^experimento-c-"}, "proyecto_id": {"$in": list(PROYECTOS)}},
        {"_id": 0, "fecha": 1, "proyecto_id": 1, "skill": 1, "referencia": 1, "modelo": 1, "resultado": 1, "duracion_ms": 1,
         "costo_usd": 1, "error": 1},
    ).sort("fecha", 1))
    for m in metricas:
        m["fecha"] = str(m["fecha"])
    (BASE / "datos" / f"metricas-{corrida}.json").write_text(json.dumps(metricas, ensure_ascii=False, indent=1), encoding="utf-8")

    resumen: dict[str, dict] = {}
    for cond in ("contexto", "solo_diff"):
        obs = [o for v in pares.values() for o in v[cond]["observaciones"]]
        resumen[cond] = {
            "prs": len(pares),
            "observaciones": len(obs),
            "prs_sin_observaciones": sum(1 for v in pares.values() if not v[cond]["observaciones"]),
            "por_severidad": dict(Counter(o["severidad"] for o in obs)),
            "por_fuente": dict(Counter(o["fuente"] for o in obs)),
            "duracion_min_total": round(sum(v[cond]["duracion_ms"] for v in pares.values()) / 60000, 1),
            "costo_nocional_usd": round(sum(m.get("costo_usd") or 0 for m in metricas if m["skill"] == f"experimento-c-{cond}"), 2),
        }
    # Llamadas que no cuentan para H1 pero sí costaron: la prueba previa y la condición «contexto» descartada (ADR-0089).
    resumen["fuera_del_experimento"] = {
        sk: {"llamadas": sum(1 for m in metricas if m["skill"] == sk),
             "costo_nocional_usd": round(sum(m.get("costo_usd") or 0 for m in metricas if m["skill"] == sk), 2)}
        for sk in sorted({m["skill"] for m in metricas} - {"experimento-c-contexto", "experimento-c-solo_diff"})
    }
    resumen["errores"] = sum(1 for m in metricas if m.get("resultado") != "ok")
    resumen["modelo"] = sorted({d["modelo"] for d in docs})
    resumen["modelos_informados_por_el_cli"] = sorted({m["modelo"] for m in metricas if m.get("modelo")})
    resumen["diffs_recortados_140k"] = sum(1 for (_, v) in pares.items() if v["contexto"]["diff_chars"] > 140_000)
    (BASE / "datos" / f"resumen-{corrida}.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=1), encoding="utf-8")

    # Paquete ciego: PRs en orden aleatorio (P01…), observaciones de las dos condiciones mezcladas dentro de cada PR.
    azar = random.Random(SEMILLA_CIEGA)
    claves = list(pares)
    azar.shuffle(claves)
    ciega = BASE / "evaluacion-ciega"
    (ciega / "contexto").mkdir(parents=True, exist_ok=True)
    clave: dict[str, dict] = {}
    filas = []
    for i, (pid, ref) in enumerate(claves, start=1):
        pr_id = f"P{i:02d}"
        v = pares[(pid, ref)]
        grupo, pt = v["contexto"]["grupo"], v["contexto"]["pt"]
        hu = db.hus.find_one({"proyecto_id": pid, "id": grupo}) if grupo != "BASE" else None
        cambio = enlace_del_cambio(v["contexto"]["pr"])
        # El paquete como lo vio la primera revisión (sin las observaciones del piloto) y sin los campos de estado,
        # que podrían sesgar a quien califica.
        md = "\n".join(l for l in md_de_la_primera_revision(pid, ref, grupo, pt).splitlines()
                       if not l.startswith(("estado:", "ronda:", "rondas_revision:")))
        partes = [f"# {pr_id} · {PROYECTOS[pid]} · {ref}", f"- Pull Request: {v['contexto']['pr']}",
                  f"- Cambio revisado (primer commit del PR, antes de correcciones): {cambio}",
                  "", "## Paquete de trabajo", "", md]
        if hu:
            partes += ["", "## Historia de usuario (spec)", "", hu.get("spec_md") or ""]
            if hu.get("test_cases_md"):
                partes += ["", "## Casos de prueba de la HU", "", hu["test_cases_md"]]
        else:
            partes += ["", "(Paquete base del proyecto: no pertenece a una HU.)"]
        (ciega / "contexto" / f"{pr_id}.md").write_text("\n".join(partes) + "\n", encoding="utf-8")

        obs = [(cond, o) for cond in ("contexto", "solo_diff") for o in v[cond]["observaciones"]]
        azar.shuffle(obs)
        for j, (cond, o) in enumerate(obs, start=1):
            oid = f"{pr_id}-{j:02d}"
            clave[oid] = {"condicion": cond, "proyecto_id": pid, "referencia": ref, "n": o["n"], "fuente": o["fuente"], "severidad": o["severidad"]}
            ubicacion = (o.get("archivo") or "") + (f":{o['linea']}" if o.get("linea") else "")
            filas.append([oid, pr_id, f"{PROYECTOS[pid]} · {ref}", ubicacion, o["descripcion"], o.get("sugerencia") or "", "", "", ""])

    wb = Workbook()
    ws = wb.active
    ws.title = "Calificación"
    encabezado = ["id", "pr", "paquete", "archivo:línea", "observación", "sugerencia", "calificación", "misma que (id)", "comentario"]
    ws.append(encabezado)
    for f in filas:
        ws.append(f)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F6F")
    anchos = [10, 6, 20, 36, 80, 60, 28, 14, 40]
    for col, ancho in zip("ABCDEFGHI", anchos):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"' + ",".join(RUBRICA) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{len(filas) + 1}")
    ws.freeze_panes = "B2"
    r = wb.create_sheet("Rúbrica")
    for fila in [["calificación", "cuándo usarla"],
                 [RUBRICA[0], "Señala un problema real del cambio, y lo que dice es correcto y accionable."],
                 [RUBRICA[1], "El problema es real, pero la explicación, la ubicación o la sugerencia son imprecisas o incompletas."],
                 [RUBRICA[2], "Es cierto pero no aporta: preferencia de estilo sin impacto, obviedad o algo que no vale la pena corregir."],
                 [RUBRICA[3], "Es incorrecta: el problema no existe o contradice el código, la HU o la arquitectura."]]:
        r.append(fila)
    r.column_dimensions["A"].width = 30
    r.column_dimensions["B"].width = 100
    wb.save(ciega / "hoja-de-evaluacion.xlsx")

    (BASE / "clave").mkdir(exist_ok=True)
    texto_clave = json.dumps({"corrida": corrida, "semilla_ciega": SEMILLA_CIEGA, "observaciones": clave}, ensure_ascii=False, indent=1)
    (BASE / "clave" / f"clave-{corrida}.json").write_text(texto_clave, encoding="utf-8")
    print(json.dumps(resumen, ensure_ascii=False, indent=1))
    print(f"observaciones a calificar: {len(filas)} en {len(claves)} PRs")
    print("SHA-256 de la clave:", hashlib.sha256(texto_clave.encode("utf-8")).hexdigest())


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "c1")
