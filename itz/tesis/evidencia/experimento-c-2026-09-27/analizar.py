"""Experimento C (ADR-0089): cruza la hoja calificada con la clave y calcula los resultados de H1.

    python analizar.py c1              calificación humana (hoja de evaluación completa) → datos/resultados-c1.json
    python analizar.py c1 opinion      calificación automática (ADR-0090 y ADR-0091)    → datos/resultados-c1-opinion.json

Los dos resultados se guardan por separado y cada uno dice quién calificó. Contenido:
- proporción de observaciones relevantes por condición (estricta: «relevante y correcta»; amplia: más «mal sustentada»);
- observaciones relevantes por PR en cada condición y la prueba pareada (Wilcoxon de rangos con signo, aproximación normal,
  con la correlación rango-biserial pareada como tamaño del efecto, y prueba de signos exacta) sobre los PRs;
- bloqueantes y mayores válidas por condición, y cuántas observaciones relevantes de una condición repite la otra (columna «misma que»)."""

import json
import math
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).parent
RELEVANTE, MAL_SUSTENTADA, RUIDO, FALSA = "relevante y correcta", "relevante pero mal sustentada", "ruido", "falsa"
CONDICIONES = ("contexto", "solo_diff")


def wilcoxon(diferencias: list[float]) -> dict:
    d = [x for x in diferencias if x != 0]
    n = len(d)
    if n == 0:
        return {"n": 0, "W+": 0, "z": 0.0, "p_bilateral": 1.0}
    orden = sorted(range(n), key=lambda i: abs(d[i]))
    rangos = [0.0] * n
    i = 0
    while i < n:
        j = i
        while j + 1 < n and abs(d[orden[j + 1]]) == abs(d[orden[i]]):
            j += 1
        for k in range(i, j + 1):
            rangos[orden[k]] = (i + j) / 2 + 1
        i = j + 1
    w_mas = sum(r for r, x in zip(rangos, d) if x > 0)
    media = n * (n + 1) / 4
    empates = Counter(abs(x) for x in d)
    var = n * (n + 1) * (2 * n + 1) / 24 - sum(t**3 - t for t in empates.values()) / 48
    z = (w_mas - media) / math.sqrt(var) if var > 0 else 0.0
    p = math.erfc(abs(z) / math.sqrt(2))
    # Tamaño del efecto: correlación rango-biserial pareada, (W+ − W−) / suma de rangos (Kerby, 2014); de −1 a 1.
    total = n * (n + 1) / 2
    r = (w_mas - (total - w_mas)) / total
    return {"n": n, "W+": w_mas, "z": round(z, 3), "p_bilateral": round(p, 4), "r_rango_biserial": round(r, 3)}


def signos(diferencias: list[float]) -> dict:
    mas, menos = sum(x > 0 for x in diferencias), sum(x < 0 for x in diferencias)
    n = mas + menos
    cola = sum(math.comb(n, k) for k in range(0, min(mas, menos) + 1)) / 2**n if n else 1.0
    return {"a_favor_contexto": mas, "a_favor_solo_diff": menos, "empates": len(diferencias) - n, "p_bilateral": round(min(1.0, 2 * cola), 4)}


def _calificacion_humana() -> tuple[dict[str, str], dict[str, str], str]:
    ws = load_workbook(BASE / "evaluacion-ciega" / "hoja-de-evaluacion.xlsx")["Calificación"]
    filas = [dict(zip([c.value for c in ws[1]], [c.value for c in fila])) for fila in ws.iter_rows(min_row=2)]
    return ({f["id"]: f["calificación"] for f in filas}, {f["id"]: (f.get("misma que (id)") or "").strip() for f in filas},
            "persona, a ciegas (hoja de evaluación)")


def _calificacion_automatica(corrida: str) -> tuple[dict[str, str], dict[str, str], str]:
    opinion = json.loads((BASE / "segunda-opinion" / f"opinion-{corrida}.json").read_text(encoding="utf-8"))
    cs = [c for pr in opinion["prs"] for c in pr["calificaciones"]]
    return ({c["id"]: c["calificacion"] for c in cs}, {c["id"]: c.get("misma_que") or "" for c in cs},
            "modelo " + ", ".join(opinion["modelos"]) + ", a ciegas (segunda opinión automática)")


def main(corrida: str, fuente: str = "hoja") -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    clave = json.loads((BASE / "clave" / f"clave-{corrida}.json").read_text(encoding="utf-8"))["observaciones"]
    cal, misma, calificador = _calificacion_automatica(corrida) if fuente == "opinion" else _calificacion_humana()
    faltan = [i for i in clave if cal.get(i) not in (RELEVANTE, MAL_SUSTENTADA, RUIDO, FALSA)]
    if faltan:
        sys.exit(f"Faltan {len(faltan)} calificaciones (p. ej. {faltan[:5]}).")
    por_cond: dict[str, Counter] = {c: Counter() for c in CONDICIONES}
    por_pr: dict[str, dict[str, dict[str, int]]] = defaultdict(lambda: {c: {"estricta": 0, "amplia": 0} for c in CONDICIONES})
    graves: dict[str, Counter] = {c: Counter() for c in CONDICIONES}
    for oid, k in clave.items():
        c, v = k["condicion"], cal[oid]
        por_cond[c][v] += 1
        pr = oid.split("-")[0]
        por_pr[pr][c]["estricta"] += v == RELEVANTE
        por_pr[pr][c]["amplia"] += v in (RELEVANTE, MAL_SUSTENTADA)
        if k["severidad"] in ("bloqueante", "mayor"):
            graves[c][f"{k['severidad']}_{'valida' if v in (RELEVANTE, MAL_SUSTENTADA) else 'no_valida'}"] += 1

    # Pares de observaciones que quien califica marcó como equivalentes («misma que»), entre condiciones distintas. La marca
    # va en una sola de las dos, así que se cuentan pares y no marcas por condición.
    pares = {tuple(sorted((oid, otra))) for oid, otra in misma.items() if otra in clave and otra != oid}
    cruzados = [p for p in pares if clave[p[0]]["condicion"] != clave[p[1]]["condicion"]]
    ambos_relevantes = [p for p in cruzados if all(cal[i] in (RELEVANTE, MAL_SUSTENTADA) for i in p)]
    con_equivalente = {i for p in cruzados for i in p}
    compartidas = Counter(clave[i]["condicion"] for i in con_equivalente if cal[i] in (RELEVANTE, MAL_SUSTENTADA))

    resultados = {"corrida": corrida, "calificador": calificador, "observaciones": len(clave), "prs": len(por_pr), "por_condicion": {}}
    for c in CONDICIONES:
        total = sum(por_cond[c].values())
        resultados["por_condicion"][c] = {
            "observaciones": total,
            "calificaciones": dict(por_cond[c]),
            "proporcion_relevantes_estricta": round(por_cond[c][RELEVANTE] / total, 3) if total else None,
            "proporcion_relevantes_amplia": round((por_cond[c][RELEVANTE] + por_cond[c][MAL_SUSTENTADA]) / total, 3) if total else None,
            "relevantes_por_pr_media_amplia": round(sum(v[c]["amplia"] for v in por_pr.values()) / len(por_pr), 2),
            "bloqueantes_y_mayores": dict(graves[c]),
            "relevantes_con_equivalente_en_la_otra_condicion": compartidas[c],
        }
    relevantes = {c: por_cond[c][RELEVANTE] + por_cond[c][MAL_SUSTENTADA] for c in CONDICIONES}
    resultados["equivalencias_entre_condiciones"] = {
        "pares_equivalentes": len(cruzados), "pares_con_ambas_relevantes": len(ambos_relevantes),
        "hallazgos_relevantes_distintos": sum(relevantes.values()) - len(ambos_relevantes),
    }
    for criterio in ("estricta", "amplia"):
        dif = [v["contexto"][criterio] - v["solo_diff"][criterio] for v in por_pr.values()]
        resultados[f"prueba_pareada_{criterio}"] = {"wilcoxon": wilcoxon(dif), "signos": signos(dif)}
    # Desglose por proyecto y por la fuente que el revisor asignó a cada observación (relevancia amplia).
    proyectos = {"9c764535fb44470a8dc92fc74c8b835e": "E1c", "b17137fd304348d7a48532ba5db89270": "E3c"}
    por_proyecto: dict[str, dict[str, list[int]]] = defaultdict(lambda: {c: [0, 0] for c in CONDICIONES})
    por_fuente: dict[str, dict[str, list[int]]] = defaultdict(lambda: {c: [0, 0] for c in CONDICIONES})
    por_severidad: dict[str, dict[str, list[int]]] = defaultdict(lambda: {c: [0, 0] for c in CONDICIONES})
    for oid, k in clave.items():
        relevante = cal[oid] in (RELEVANTE, MAL_SUSTENTADA)
        for tabla, llave in ((por_proyecto, proyectos.get(k["proyecto_id"], k["proyecto_id"])), (por_fuente, k["fuente"]), (por_severidad, k["severidad"])):
            tabla[llave][k["condicion"]][0] += relevante
            tabla[llave][k["condicion"]][1] += 1
    resultados["relevantes_y_total_por_proyecto"] = por_proyecto
    resultados["relevantes_y_total_por_fuente"] = por_fuente
    resultados["relevantes_y_total_por_severidad"] = por_severidad
    resultados["por_pr"] = dict(sorted(por_pr.items()))
    sufijo = "-opinion" if fuente == "opinion" else ""
    (BASE / "datos" / f"resultados-{corrida}{sufijo}.json").write_text(json.dumps(resultados, ensure_ascii=False, indent=1), encoding="utf-8")
    print(json.dumps({k: v for k, v in resultados.items() if k != "por_pr"}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "c1", sys.argv[2] if len(sys.argv) > 2 else "hoja")
